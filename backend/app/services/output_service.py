import io
import os
import struct
import uuid
import zlib
from datetime import datetime
from typing import Dict, List, Optional, Tuple

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.agents.output_agent import OutputAgent
from app.core.exceptions import SessionNotFoundError, StorageError
from app.models.db.models import (
    Presentation,
    QueryExecution,
    Report,
    Session as SessionModel,
    Visualization,
)

_OUTPUT_DIR = "/tmp/aria_outputs"

_MIME_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def _ensure_output_dir() -> None:
    os.makedirs(_OUTPUT_DIR, exist_ok=True)


# ── Chart rendering ────────────────────────────────────────────────────────────

def _placeholder_png(width: int = 900, height: int = 500) -> bytes:
    """Minimal valid PNG — solid light-grey rectangle."""
    r, g, b = 224, 231, 240
    raw_row = b"\x00" + bytes([r, g, b] * width)
    raw_data = raw_row * height
    compressed = zlib.compress(raw_data, 6)

    def _chunk(tag: bytes, data: bytes) -> bytes:
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    png = b"\x89PNG\r\n\x1a\n"
    png += _chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
    png += _chunk(b"IDAT", compressed)
    png += _chunk(b"IEND", b"")
    return png


def _render_chart_png(viz: Visualization, query_exec: Optional[QueryExecution]) -> bytes:
    """Build a chart PNG from visualization + query data; placeholder on any failure."""
    try:
        import pandas as pd
        import plotly.graph_objects as go

        cfg: dict = viz.chart_config or {}
        title = cfg.get("title") or viz.title or "Chart"
        x_col: str = cfg.get("x_axis", "")
        y_col: str = cfg.get("y_axis", "")
        color_col: Optional[str] = cfg.get("color_by")
        chart_type: str = viz.chart_type or "bar"

        df = pd.DataFrame()
        if query_exec and query_exec.result_preview:
            cols = query_exec.result_preview.get("columns", [])
            rows = query_exec.result_preview.get("rows", [])
            if cols and rows:
                df = pd.DataFrame(rows, columns=cols)

        if df.empty or x_col not in df.columns or y_col not in df.columns:
            raise ValueError("insufficient data for chart")

        if chart_type == "line":
            fig = go.Figure(go.Scatter(x=df[x_col], y=df[y_col], mode="lines+markers"))
        elif chart_type == "pie":
            fig = go.Figure(go.Pie(labels=df[x_col], values=df[y_col]))
        elif chart_type == "scatter":
            color = df[color_col] if color_col and color_col in df.columns else None
            fig = go.Figure(go.Scatter(x=df[x_col], y=df[y_col], mode="markers",
                                       marker=dict(color=color)))
        else:
            fig = go.Figure(go.Bar(x=df[x_col], y=df[y_col]))

        fig.update_layout(title=title, width=900, height=500,
                          margin=dict(l=40, r=40, t=60, b=40))
        return fig.to_image(format="png")
    except Exception:
        return _placeholder_png()


# ── Document builders ──────────────────────────────────────────────────────────

def _build_docx(plan: dict, chart_images: Dict[str, bytes]) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    title_para = doc.add_heading(plan.get("title", "Analysis Report"), level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    date_para = doc.add_paragraph(f"Generated: {datetime.utcnow().strftime('%B %d, %Y')}")
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(plan.get("executive_summary", ""))

    for section in plan.get("sections", []):
        doc.add_heading(section.get("heading", ""), level=2)
        doc.add_paragraph(section.get("content", ""))

        points: List[str] = section.get("talking_points", [])
        if points:
            doc.add_heading("Key Points", level=3)
            for pt in points:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(pt)

        viz_id = section.get("visualization_id")
        if viz_id and viz_id in chart_images:
            doc.add_picture(io.BytesIO(chart_images[viz_id]), width=Inches(6))

        doc.add_paragraph()

    recs: List[str] = plan.get("recommendations", [])
    if recs:
        doc.add_heading("Recommendations", level=1)
        for rec in recs:
            p = doc.add_paragraph(style="List Number")
            p.add_run(rec)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_PDF_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<style>
  body{font-family:Arial,sans-serif;margin:48px;color:#222;font-size:13px}
  h1{font-size:26px;color:#1a1a2e;border-bottom:2px solid #4a90e2;padding-bottom:8px;text-align:center}
  h2{font-size:18px;color:#2d4a7a;margin-top:28px}
  h3{font-size:14px;color:#4a6fa5}
  .date{text-align:center;color:#888;margin-bottom:24px}
  .exec{background:#f0f4ff;padding:18px;border-radius:6px;margin:16px 0}
  .section{margin-bottom:28px;page-break-inside:avoid}
  ul{margin:6px 0 12px 20px} li{margin:4px 0}
  .chart{text-align:center;margin:16px 0}
  .chart img{max-width:100%;border:1px solid #ddd;border-radius:4px}
  .recs{background:#f9f9f9;padding:18px;border-radius:6px}
  ol li{margin:8px 0}
</style>
</head>
<body>
<h1>{{ title }}</h1>
<p class="date">Generated: {{ date }}</p>
<div class="exec"><h2>Executive Summary</h2><p>{{ executive_summary }}</p></div>
{% for s in sections %}
<div class="section">
  <h2>{{ s.heading }}</h2>
  <p>{{ s.content }}</p>
  {% if s.talking_points %}
  <h3>Key Points</h3><ul>{% for p in s.talking_points %}<li>{{ p }}</li>{% endfor %}</ul>
  {% endif %}
  {% if s.visualization_id and s.visualization_id in chart_b64 %}
  <div class="chart">
    <img src="data:image/png;base64,{{ chart_b64[s.visualization_id] }}" alt="Chart">
  </div>
  {% endif %}
</div>
{% endfor %}
{% if recommendations %}
<div class="recs"><h2>Recommendations</h2>
<ol>{% for r in recommendations %}<li>{{ r }}</li>{% endfor %}</ol>
</div>
{% endif %}
</body>
</html>"""


def _build_pdf(plan: dict, chart_images: Dict[str, bytes]) -> bytes:
    import base64

    from jinja2 import BaseLoader, Environment

    chart_b64 = {vid: base64.b64encode(img).decode() for vid, img in chart_images.items()}
    env = Environment(loader=BaseLoader(), autoescape=True)
    html = env.from_string(_PDF_TEMPLATE).render(
        title=plan.get("title", "Analysis Report"),
        date=datetime.utcnow().strftime("%B %d, %Y"),
        executive_summary=plan.get("executive_summary", ""),
        sections=plan.get("sections", []),
        recommendations=plan.get("recommendations", []),
        chart_b64=chart_b64,
    )

    try:
        import weasyprint
        return weasyprint.HTML(string=html).write_pdf()
    except ImportError as exc:
        raise StorageError(f"PDF generation requires WeasyPrint: {exc}") from exc


def _build_pptx(plan: dict, chart_images: Dict[str, bytes]) -> bytes:
    from pptx import Presentation as PptxPresentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = PptxPresentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]
    title_layout = prs.slide_layouts[0]

    # Title slide
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = plan.get("title", "Analysis Report")
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = datetime.utcnow().strftime("Generated %B %d, %Y")

    def _heading(sl, text: str) -> None:
        tb = sl.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12), Inches(1))
        tf = tb.text_frame
        tf.word_wrap = True
        run = tf.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    def _body(sl, text: str, top=None, width=None) -> None:
        top = top or Inches(1.4)
        width = width or Inches(12.3)
        tb = sl.shapes.add_textbox(Inches(0.5), top, width, Inches(5.6))
        tf = tb.text_frame
        tf.word_wrap = True
        run = tf.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(14)

    # Executive summary slide
    slide = prs.slides.add_slide(blank)
    _heading(slide, "Executive Summary")
    _body(slide, plan.get("executive_summary", ""))

    # Section slides
    for section in plan.get("sections", []):
        slide = prs.slides.add_slide(blank)
        _heading(slide, section.get("heading", ""))

        viz_id = section.get("visualization_id")
        has_chart = bool(viz_id and viz_id in chart_images)

        content = section.get("content", "")
        points: List[str] = section.get("talking_points", [])
        body_text = content
        if points:
            bullet_block = "\n".join(f"  • {p}" for p in points)
            body_text = f"{content}\n\n{bullet_block}" if content else bullet_block

        body_width = Inches(6.5) if has_chart else Inches(12.3)
        _body(slide, body_text, width=body_width)

        if has_chart:
            slide.shapes.add_picture(
                io.BytesIO(chart_images[viz_id]),
                Inches(7.1), Inches(1.4), Inches(5.8), Inches(5.0),
            )

    # Recommendations slide
    recs: List[str] = plan.get("recommendations", [])
    if recs:
        slide = prs.slides.add_slide(blank)
        _heading(slide, "Recommendations")
        rec_text = "\n".join(f"{i}. {r}" for i, r in enumerate(recs, 1))
        _body(slide, rec_text)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ── Service class ──────────────────────────────────────────────────────────────

class OutputService:
    def __init__(self, db: AsyncSession):
        self.db = db

    # ── Reports ────────────────────────────────────────────────────────────────

    async def generate_report(
        self,
        session_id: uuid.UUID,
        report_format: str,
        audience: Optional[str],
        focus: Optional[str],
        report_id: Optional[uuid.UUID] = None,
    ) -> Report:
        report = await self._load_or_create_report(
            report_id=report_id,
            session_id=session_id,
            report_format=report_format,
        )
        report.status = "generating"
        await self.db.flush()

        try:
            agent = OutputAgent(db=self.db)
            plan = await agent.run(
                input={
                    "session_id": str(session_id),
                    "output_type": "report",
                    "audience": audience,
                    "focus": focus,
                },
                session_context={},
            )

            chart_images = await self._collect_chart_images(plan)

            if report_format == "pdf":
                file_bytes = _build_pdf(plan, chart_images)
            else:
                file_bytes = _build_docx(plan, chart_images)

            storage_key = self._save_file(file_bytes, report_format)
            report.title = plan.get("title", report.title)
            report.storage_key = storage_key
            report.status = "ready"
            await self.db.flush()
        except Exception as exc:
            report.status = "failed"
            await self.db.flush()
            raise StorageError(f"Report generation failed: {exc}") from exc

        return report

    async def get_report(self, report_id: uuid.UUID) -> Optional[Report]:
        result = await self.db.execute(select(Report).where(Report.id == report_id))
        return result.scalar_one_or_none()

    async def list_reports(self, session_id: uuid.UUID) -> List[Report]:
        result = await self.db.execute(
            select(Report).where(Report.session_id == session_id).order_by(Report.created_at.desc())
        )
        return list(result.scalars().all())

    # ── Presentations ──────────────────────────────────────────────────────────

    async def generate_presentation(
        self,
        session_id: uuid.UUID,
        pres_format: str,
        audience: Optional[str],
        focus: Optional[str],
        presentation_id: Optional[uuid.UUID] = None,
    ) -> Presentation:
        pres = await self._load_or_create_presentation(
            presentation_id=presentation_id,
            session_id=session_id,
            pres_format=pres_format,
        )
        pres.status = "generating"
        await self.db.flush()

        try:
            agent = OutputAgent(db=self.db)
            plan = await agent.run(
                input={
                    "session_id": str(session_id),
                    "output_type": "presentation",
                    "audience": audience,
                    "focus": focus,
                },
                session_context={},
            )

            chart_images = await self._collect_chart_images(plan)
            file_bytes = _build_pptx(plan, chart_images)

            storage_key = self._save_file(file_bytes, pres_format)
            pres.title = plan.get("title", pres.title)
            pres.storage_key = storage_key
            pres.status = "ready"
            await self.db.flush()
        except Exception as exc:
            pres.status = "failed"
            await self.db.flush()
            raise StorageError(f"Presentation generation failed: {exc}") from exc

        return pres

    async def get_presentation(self, presentation_id: uuid.UUID) -> Optional[Presentation]:
        result = await self.db.execute(
            select(Presentation).where(Presentation.id == presentation_id)
        )
        return result.scalar_one_or_none()

    async def list_presentations(self, session_id: uuid.UUID) -> List[Presentation]:
        result = await self.db.execute(
            select(Presentation)
            .where(Presentation.session_id == session_id)
            .order_by(Presentation.created_at.desc())
        )
        return list(result.scalars().all())

    # ── File download ──────────────────────────────────────────────────────────

    async def download_file(self, storage_key: str) -> Tuple[bytes, str]:
        if not os.path.exists(storage_key):
            raise StorageError(f"File not found: {storage_key}")
        ext = storage_key.rsplit(".", 1)[-1].lower()
        mime = _MIME_TYPES.get(ext, "application/octet-stream")
        with open(storage_key, "rb") as f:
            return f.read(), mime

    # ── Private helpers ────────────────────────────────────────────────────────

    async def _load_or_create_report(
        self,
        report_id: Optional[uuid.UUID],
        session_id: uuid.UUID,
        report_format: str,
    ) -> Report:
        if report_id:
            result = await self.db.execute(select(Report).where(Report.id == report_id))
            report = result.scalar_one_or_none()
            if report:
                return report

        session = await self._load_session_title(session_id)
        report = Report(
            session_id=session_id,
            title=session or "Analysis Report",
            report_format=report_format,
            status="pending",
        )
        self.db.add(report)
        await self.db.flush()
        await self.db.refresh(report)
        return report

    async def _load_or_create_presentation(
        self,
        presentation_id: Optional[uuid.UUID],
        session_id: uuid.UUID,
        pres_format: str,
    ) -> Presentation:
        if presentation_id:
            result = await self.db.execute(
                select(Presentation).where(Presentation.id == presentation_id)
            )
            pres = result.scalar_one_or_none()
            if pres:
                return pres

        session = await self._load_session_title(session_id)
        pres = Presentation(
            session_id=session_id,
            title=session or "Analysis Presentation",
            pres_format=pres_format,
            status="pending",
        )
        self.db.add(pres)
        await self.db.flush()
        await self.db.refresh(pres)
        return pres

    async def _load_session_title(self, session_id: uuid.UUID) -> str:
        result = await self.db.execute(
            select(SessionModel.title).where(SessionModel.id == session_id)
        )
        return result.scalar_one_or_none() or "Analysis Report"

    async def _collect_chart_images(self, plan: dict) -> Dict[str, bytes]:
        viz_ids = {
            s["visualization_id"]
            for s in plan.get("sections", [])
            if s.get("visualization_id")
        }
        if not viz_ids:
            return {}

        uuid_ids = []
        for vid in viz_ids:
            try:
                uuid_ids.append(uuid.UUID(vid))
            except ValueError:
                pass

        if not uuid_ids:
            return {}

        viz_result = await self.db.execute(
            select(Visualization).where(Visualization.id.in_(uuid_ids))
        )
        visualizations = viz_result.scalars().all()

        exec_ids = [v.query_execution_id for v in visualizations if v.query_execution_id]
        exec_map: Dict[uuid.UUID, QueryExecution] = {}
        if exec_ids:
            exec_result = await self.db.execute(
                select(QueryExecution).where(QueryExecution.id.in_(exec_ids))
            )
            for qe in exec_result.scalars().all():
                exec_map[qe.id] = qe

        images: Dict[str, bytes] = {}
        for viz in visualizations:
            qe = exec_map.get(viz.query_execution_id) if viz.query_execution_id else None
            images[str(viz.id)] = _render_chart_png(viz, qe)

        return images

    @staticmethod
    def _save_file(file_bytes: bytes, ext: str) -> str:
        _ensure_output_dir()
        filename = f"{uuid.uuid4()}.{ext}"
        path = os.path.join(_OUTPUT_DIR, filename)
        with open(path, "wb") as f:
            f.write(file_bytes)
        return path
